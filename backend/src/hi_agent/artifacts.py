"""Tenant-safe generated artifact storage and local document renderers."""

from __future__ import annotations

import base64
import hashlib
import io
import re
from pathlib import Path
from typing import Any, Literal

import httpx
from docx import Document as WordDocument
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PptPt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from sqlalchemy.orm import Session

from .config import Settings
from .errors import HiAgentError
from .models import Artifact, ImageEndpoint


def _safe_filename(value: str, suffix: str) -> str:
    stem = re.sub(r"[^\w\u3400-\u9fff.-]+", "-", value, flags=re.UNICODE).strip(".-")[:100] or "artifact"
    return f"{stem}.{suffix}"


class ArtifactService:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def create_report(
        self,
        db: Session,
        *,
        owner_id: str,
        project_id: str,
        title: str,
        content: str,
        output_format: Literal["md", "docx", "pdf"],
        run_id: str | None = None,
    ) -> Artifact:
        if output_format == "md":
            payload = f"# {title}\n\n{content}\n".encode()
            media_type = "text/markdown; charset=utf-8"
        elif output_format == "docx":
            document = WordDocument()
            document.add_heading(title, 0)
            for block in content.split("\n"):
                paragraph = document.add_paragraph(block)
                if paragraph.style is not None:
                    paragraph.style.font.size = Pt(11)
            buffer = io.BytesIO()
            document.save(buffer)
            payload = buffer.getvalue()
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            buffer = io.BytesIO()
            doc = SimpleDocTemplate(buffer, pagesize=A4, title=title)
            styles = getSampleStyleSheet()
            story = [Paragraph(title, styles["Title"]), Spacer(1, 16)]
            for block in content.split("\n"):
                if block.strip():
                    story.extend(
                        [Paragraph(block.replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"]), Spacer(1, 6)]
                    )
            doc.build(story)
            payload = buffer.getvalue()
            media_type = "application/pdf"
        return self._store(
            db, owner_id, project_id, run_id, "report", _safe_filename(title, output_format), media_type, payload
        )

    def create_presentation(
        self,
        db: Session,
        *,
        owner_id: str,
        project_id: str,
        title: str,
        slides: list[dict[str, Any]],
        run_id: str | None = None,
    ) -> Artifact:
        if not slides or len(slides) > self.settings.artifact_max_slides:
            raise HiAgentError("ARTIFACT_SLIDE_LIMIT", "幻灯片页数超出限制")
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        cover = presentation.slides.add_slide(presentation.slide_layouts[0])
        cover.shapes.title.text = title
        cover.placeholders[1].text = "由 Hi-agent 生成"
        for raw in slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = str(raw.get("title", "内容"))[:200]
            frame = slide.placeholders[1].text_frame
            frame.clear()
            bullets = raw.get("bullets", [])
            if not isinstance(bullets, list):
                bullets = [str(bullets)]
            for index, bullet in enumerate(bullets[:10]):
                paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                paragraph.text = str(bullet)[:1000]
                paragraph.font.size = PptPt(24)
        buffer = io.BytesIO()
        presentation.save(buffer)
        return self._store(
            db,
            owner_id,
            project_id,
            run_id,
            "presentation",
            _safe_filename(title, "pptx"),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            buffer.getvalue(),
        )

    async def create_image(
        self,
        db: Session,
        *,
        owner_id: str,
        project_id: str,
        prompt: str,
        endpoint: ImageEndpoint,
        run_id: str | None = None,
    ) -> Artifact:
        api_key = self.settings.resolve_secret(endpoint.api_key_env)
        if not api_key:
            raise HiAgentError("IMAGE_API_KEY_MISSING", "图像端点引用的密钥环境变量未配置")
        url = endpoint.base_url.rstrip("/")
        if not url.endswith("/images/generations"):
            url = f"{url}/images/generations"
        try:
            async with httpx.AsyncClient(timeout=endpoint.timeout_seconds, follow_redirects=False) as client:
                response = await client.post(
                    url,
                    headers={"Authorization": f"Bearer {api_key}"},
                    json={"model": endpoint.model, "prompt": prompt[:8000], "n": 1, "response_format": "b64_json"},
                )
            response.raise_for_status()
            data = response.json().get("data", [])
            encoded = data[0].get("b64_json") if data and isinstance(data[0], dict) else None
            if not encoded:
                raise ValueError("endpoint did not return b64_json")
            payload = base64.b64decode(encoded, validate=True)
            if payload.startswith(b"\x89PNG\r\n\x1a\n"):
                media_type, suffix = "image/png", "png"
            elif payload.startswith(b"\xff\xd8\xff"):
                media_type, suffix = "image/jpeg", "jpg"
            elif payload.startswith(b"RIFF") and payload[8:12] == b"WEBP":
                media_type, suffix = "image/webp", "webp"
            else:
                raise ValueError("endpoint returned an unsupported image format")
        except (httpx.HTTPError, ValueError, KeyError) as exc:
            raise HiAgentError("IMAGE_GENERATION_FAILED", "图像生成端点调用失败", details={"reason": str(exc)}) from exc
        return self._store(
            db,
            owner_id,
            project_id,
            run_id,
            "image",
            _safe_filename("generated-image", suffix),
            media_type,
            payload,
        )

    def _store(
        self,
        db: Session,
        owner_id: str,
        project_id: str,
        run_id: str | None,
        kind: str,
        filename: str,
        media_type: str,
        payload: bytes,
    ) -> Artifact:
        if len(payload) > self.settings.artifact_max_bytes:
            raise HiAgentError("ARTIFACT_TOO_LARGE", "生成文件超过大小限制", status_code=413)
        artifact = Artifact(
            owner_id=owner_id,
            project_id=project_id,
            run_id=run_id,
            kind=kind,
            filename=filename,
            media_type=media_type,
            size_bytes=len(payload),
            sha256=hashlib.sha256(payload).hexdigest(),
            storage_path="",
        )
        db.add(artifact)
        db.flush()
        root = self.settings.artifacts_dir.resolve()
        target_dir = (root / owner_id / project_id / artifact.id).resolve()
        if not target_dir.is_relative_to(root):
            raise HiAgentError("ARTIFACT_PATH_ESCAPE", "生成文件路径越界")
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / filename
        temporary = target.with_suffix(target.suffix + ".tmp")
        temporary.write_bytes(payload)
        temporary.replace(target)
        artifact.storage_path = str(target)
        db.commit()
        db.refresh(artifact)
        return artifact

    def resolve_download(self, artifact: Artifact) -> Path:
        root = self.settings.artifacts_dir.resolve()
        path = Path(artifact.storage_path).resolve()
        if not path.is_relative_to(root) or not path.is_file() or path.is_symlink():
            raise HiAgentError("ARTIFACT_FILE_MISSING", "生成文件不存在或路径无效", status_code=404)
        return path
