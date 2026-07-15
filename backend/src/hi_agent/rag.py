"""Document parsing, token-aware chunking and local vector retrieval."""

from __future__ import annotations

import hashlib
import io
import math
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from .config import Settings, get_settings
from .errors import ConflictError, HiAgentError, NotFoundError, ServiceUnavailableError
from .models import Document, DocumentChunk, KnowledgeBase
from .schemas import SearchHit

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".markdown", ".txt"}
_TOKEN = re.compile(r"[\u3400-\u9fff]|[A-Za-z0-9_]+|[^\s]", re.UNICODE)


@dataclass(frozen=True)
class ParsedPage:
    text: str
    page: int | None


@dataclass(frozen=True)
class Chunk:
    content: str
    page: int | None
    chunk_index: int


def validate_filename(filename: str) -> str:
    if not filename or filename != Path(filename).name or "\x00" in filename:
        raise HiAgentError("INVALID_FILENAME", "文件名无效")
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise HiAgentError(
            "UNSUPPORTED_DOCUMENT_TYPE",
            "仅支持 PDF、DOCX、Markdown 和 TXT",
            details={"extension": suffix},
        )
    return filename


def parse_document(filename: str, content: bytes) -> list[ParsedPage]:
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            pages = [ParsedPage(page.extract_text() or "", index + 1) for index, page in enumerate(reader.pages)]
        elif suffix == ".docx":
            from docx import Document as DocxDocument

            document = DocxDocument(io.BytesIO(content))
            pages = [ParsedPage("\n".join(p.text for p in document.paragraphs), None)]
        else:
            pages = [ParsedPage(content.decode("utf-8-sig"), None)]
    except (UnicodeDecodeError, ValueError, OSError) as exc:
        raise HiAgentError("DOCUMENT_PARSE_FAILED", "文档解析失败", details={"reason": str(exc)}) from exc
    except Exception as exc:
        # Parser libraries expose a broad set of format-specific exceptions.
        raise HiAgentError("DOCUMENT_PARSE_FAILED", "文档损坏或无法解析", details={"reason": str(exc)}) from exc
    if not any(page.text.strip() for page in pages):
        raise HiAgentError("EMPTY_DOCUMENT", "文档中没有可索引文本")
    return pages


def chunk_pages(pages: list[ParsedPage], chunk_size: int, overlap: int) -> list[Chunk]:
    chunks: list[Chunk] = []
    next_index = 0
    step = chunk_size - overlap
    if step <= 0:
        raise ValueError("overlap must be smaller than chunk_size")
    for page in pages:
        text = page.text.strip()
        matches = list(_TOKEN.finditer(text))
        if not matches:
            continue
        start = 0
        while start < len(matches):
            end = min(start + chunk_size, len(matches))
            char_start = matches[start].start()
            char_end = matches[end - 1].end()
            chunk_text = text[char_start:char_end].strip()
            if chunk_text:
                chunks.append(Chunk(content=chunk_text, page=page.page, chunk_index=next_index))
                next_index += 1
            if end == len(matches):
                break
            start += step
    return chunks


class EmbeddingProvider:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings
        self._models: dict[str, Any] = {}

    @staticmethod
    def _deterministic(text: str, dimensions: int = 96) -> list[float]:
        vector = [0.0] * dimensions
        tokens = [match.group(0).lower() for match in _TOKEN.finditer(text)]
        for token in tokens:
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            slot = int.from_bytes(digest[:4], "big") % dimensions
            sign = 1.0 if digest[4] % 2 else -1.0
            vector[slot] += sign
        norm = math.sqrt(sum(item * item for item in vector)) or 1.0
        return [item / norm for item in vector]

    def embed(self, texts: list[str], model_name: str) -> list[list[float]]:
        if self.settings.embedding_backend == "deterministic":
            return [self._deterministic(text) for text in texts]
        try:
            from fastembed import TextEmbedding

            model = self._models.get(model_name)
            if model is None:
                model = TextEmbedding(model_name=model_name, cache_dir=str(self.settings.data_dir / "models"))
                self._models[model_name] = model
            return [embedding.tolist() for embedding in model.embed(texts)]
        except Exception as exc:
            raise ServiceUnavailableError(
                "EMBEDDING_UNAVAILABLE",
                "嵌入模型不可用，未切换到外部服务",
                reason=str(exc),
                model=model_name,
            ) from exc


class RagService:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()
        self.embeddings = EmbeddingProvider(self.settings)
        self._qdrant: Any | None = None

    def _client(self) -> Any:
        if self.settings.embedding_backend == "deterministic":
            return None
        if self._qdrant is None:
            try:
                from qdrant_client import QdrantClient

                self.settings.vectors_dir.mkdir(parents=True, exist_ok=True)
                self._qdrant = QdrantClient(path=str(self.settings.vectors_dir))
            except Exception as exc:
                raise ServiceUnavailableError(
                    "VECTOR_STORE_UNAVAILABLE", "Qdrant 本地向量库不可用", reason=str(exc)
                ) from exc
        return self._qdrant

    @staticmethod
    def _collection(kb_id: str) -> str:
        return f"kb_{kb_id.replace('-', '_')}"

    def ingest(self, db: Session, kb: KnowledgeBase, filename: str, media_type: str, content: bytes) -> Document:
        filename = validate_filename(filename)
        digest = hashlib.sha256(content).hexdigest()
        duplicate = db.scalar(
            select(Document).where(
                Document.knowledge_base_id == kb.id,
                Document.sha256 == digest,
            )
        )
        if duplicate is not None:
            raise ConflictError(
                "DOCUMENT_DUPLICATE",
                "该知识库已包含相同内容的文档",
                document_id=duplicate.id,
            )
        document = Document(
            knowledge_base_id=kb.id,
            filename=filename,
            media_type=media_type or "application/octet-stream",
            sha256=digest,
            size_bytes=len(content),
            storage_path="",
        )
        db.add(document)
        db.flush()
        suffix = Path(filename).suffix.lower()
        target_dir = (self.settings.uploads_dir / kb.id).resolve()
        target_dir.mkdir(parents=True, exist_ok=True)
        target = (target_dir / f"{document.id}{suffix}").resolve()
        if not target.is_relative_to(self.settings.uploads_dir.resolve()):
            raise HiAgentError("UPLOAD_PATH_ESCAPE", "上传路径越界")
        document.storage_path = str(target)
        target.write_bytes(content)
        try:
            parsed = parse_document(filename, content)
            chunks = chunk_pages(parsed, kb.chunk_size, kb.chunk_overlap)
            vectors = self.embeddings.embed([chunk.content for chunk in chunks], kb.embedding_model)
            for chunk, vector in zip(chunks, vectors, strict=True):
                db.add(
                    DocumentChunk(
                        document_id=document.id,
                        chunk_index=chunk.chunk_index,
                        page=chunk.page,
                        content=chunk.content,
                        vector_json=vector,
                    )
                )
            if self.settings.embedding_backend != "deterministic":
                self._upsert_qdrant(kb, document, chunks, vectors)
            document.status = "ready"
            document.chunk_count = len(chunks)
            db.commit()
            db.refresh(document)
            return document
        except Exception as exc:
            document.status = "failed"
            document.error = str(exc)
            db.commit()
            raise

    def _upsert_qdrant(
        self,
        kb: KnowledgeBase,
        document: Document,
        chunks: list[Chunk],
        vectors: list[list[float]],
    ) -> None:
        if not vectors:
            return
        from qdrant_client import models as qmodels

        client = self._client()
        collection = self._collection(kb.id)
        if not client.collection_exists(collection):
            client.create_collection(
                collection_name=collection,
                vectors_config=qmodels.VectorParams(size=len(vectors[0]), distance=qmodels.Distance.COSINE),
            )
        points = [
            qmodels.PointStruct(
                id=str(uuid.uuid5(uuid.UUID(document.id), str(chunk.chunk_index))),
                vector=vector,
                payload={
                    "document_id": document.id,
                    "filename": document.filename,
                    "page": chunk.page,
                    "chunk_index": chunk.chunk_index,
                    "content": chunk.content,
                },
            )
            for chunk, vector in zip(chunks, vectors, strict=True)
        ]
        client.upsert(collection_name=collection, points=points, wait=True)

    def search(self, db: Session, kb: KnowledgeBase, query: str, top_k: int | None = None) -> list[SearchHit]:
        limit = top_k or kb.top_k
        vector = self.embeddings.embed([query], kb.embedding_model)[0]
        if self.settings.embedding_backend == "deterministic":
            return self._search_database(db, kb.id, vector, limit)
        client = self._client()
        collection = self._collection(kb.id)
        if not client.collection_exists(collection):
            return []
        try:
            response = client.query_points(collection_name=collection, query=vector, limit=limit)
            points = response.points
        except AttributeError:
            points = client.search(collection_name=collection, query_vector=vector, limit=limit)
        return [
            SearchHit(
                document_id=str(point.payload["document_id"]),
                filename=str(point.payload["filename"]),
                page=point.payload.get("page"),
                chunk_index=int(point.payload["chunk_index"]),
                content=str(point.payload["content"]),
                score=float(point.score),
            )
            for point in points
        ]

    @staticmethod
    def _search_database(db: Session, kb_id: str, query: list[float], limit: int) -> list[SearchHit]:
        rows = db.execute(
            select(DocumentChunk, Document)
            .join(Document, Document.id == DocumentChunk.document_id)
            .where(Document.knowledge_base_id == kb_id, Document.status == "ready")
        ).all()
        scored: list[tuple[float, DocumentChunk, Document]] = []
        for chunk, document in rows:
            vector = chunk.vector_json or []
            score = sum(left * right for left, right in zip(query, vector, strict=False))
            scored.append((score, chunk, document))
        scored.sort(key=lambda item: item[0], reverse=True)
        return [
            SearchHit(
                document_id=document.id,
                filename=document.filename,
                page=chunk.page,
                chunk_index=chunk.chunk_index,
                content=chunk.content,
                score=float(score),
            )
            for score, chunk, document in scored[:limit]
        ]

    def delete_document(self, db: Session, document: Document) -> None:
        kb_id = document.knowledge_base_id
        if self.settings.embedding_backend != "deterministic":
            client = self._client()
            collection = self._collection(kb_id)
            if client.collection_exists(collection):
                from qdrant_client import models as qmodels

                client.delete(
                    collection_name=collection,
                    points_selector=qmodels.FilterSelector(
                        filter=qmodels.Filter(
                            must=[
                                qmodels.FieldCondition(
                                    key="document_id", match=qmodels.MatchValue(value=document.id)
                                )
                            ]
                        )
                    ),
                    wait=True,
                )
        path = Path(document.storage_path)
        if path.exists() and path.resolve().is_relative_to(self.settings.uploads_dir.resolve()):
            path.unlink()
        db.delete(document)
        db.commit()

    @staticmethod
    def get_knowledge_base(db: Session, kb_id: str) -> KnowledgeBase:
        kb = db.get(KnowledgeBase, kb_id)
        if kb is None:
            raise NotFoundError("KnowledgeBase", kb_id)
        return kb


_rag_service: RagService | None = None


def get_rag_service() -> RagService:
    global _rag_service
    if _rag_service is None:
        _rag_service = RagService()
    return _rag_service


def reset_rag_service() -> None:
    global _rag_service
    _rag_service = None

